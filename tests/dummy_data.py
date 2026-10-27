"""テスト用のダミー資料（PDF・動画・画像・Excel・ZIP）を生成する。個人情報・実データは含まない。"""
from __future__ import annotations

import io
import os
import subprocess
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

JP_FONTS = ["/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf",
            "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"]


def _font(size: int):
    for f in JP_FONTS:
        if os.path.exists(f):
            return ImageFont.truetype(f, size)
    try:
        return ImageFont.truetype("DejaVuSans.ttf", size)
    except OSError:
        return ImageFont.load_default()


def text_image(path: Path, lines: list[str], size=(1280, 720), bg="white", fg="black", font_size=48) -> Path:
    im = Image.new("RGB", size, bg)
    d = ImageDraw.Draw(im)
    f = _font(font_size)
    y = 40
    for ln in lines:
        d.text((40, y), ln, fill=fg, font=f)
        y += int(font_size * 1.5)
    # 表っぽい枠線
    d.rectangle([40, y + 10, 640, y + 130], outline=fg, width=3)
    d.line([340, y + 10, 340, y + 130], fill=fg, width=3)
    d.line([40, y + 70, 640, y + 70], fill=fg, width=3)
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path)
    return path


def make_pdf(path: Path, deadline: str = "11月20日", pages_extra: int = 0) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4

    def page(lines: list[str]):
        c.setFont("HeiseiKakuGo-W5", 12)
        y = h - 60
        for ln in lines:
            c.drawString(50, y, ln)
            y -= 22
        return y

    # p.1 目的・手順
    page(["鋼材の引張試験（ダミー資料）",
          "実験の目的：鋼材の応力度－ひずみ度関係を求め、ヤング係数と降伏点を確認する。",
          "手順1：試験体の断面寸法をノギスで測定し記録すること。",
          "手順2：荷重と変位を一定間隔で記録してください。",
          "注意：試験機の周囲に立ち入ってはいけない。"])
    c.showPage()
    # p.2 数式・表・グラフ（ベクター）
    y = page(["応力度の計算式： σ = P / A",
              "ひずみ度： ε = ΔL / L",
              "ヤング係数 E = σ / ε （弾性域の傾きから求める）",
              f"レポートの提出期限は{deadline}です。PDF形式で提出してください。"])
    c.setFont("HeiseiKakuGo-W5", 10)
    rows = [["荷重P(kN)", "変位(mm)"], ["10", "0.12"], ["20", "0.25"], ["30", "0.37"]]
    top = y - 20
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c.rect(50 + j * 100, top - i * 20, 100, 20)
            c.drawString(55 + j * 100, top - i * 20 + 6, v)
    gy = top - 140
    c.line(50, gy, 300, gy)
    c.line(50, gy, 50, gy + 120)
    pts = [(50 + k * 25, gy + k * 11) for k in range(10)]
    for a, b in zip(pts, pts[1:]):
        c.line(*a, *b)
    for k in range(25):  # 目盛り（ベクター図形要素）
        c.line(50 + k * 10, gy - 3, 50 + k * 10, gy + 3)
    c.showPage()
    # p.3 スキャン画像のみ（テキストなし）
    buf = io.BytesIO()
    tmp = Image.new("RGB", (1200, 500), "white")
    ImageDraw.Draw(tmp).text((30, 30), "Scanned page: Yield point must be marked", fill="black", font=_font(40))
    tmp.save(buf, format="PNG")
    from reportlab.lib.utils import ImageReader
    c.drawImage(ImageReader(io.BytesIO(buf.getvalue())), 40, h - 400, width=500, height=200)
    c.showPage()
    for k in range(pages_extra):
        page([f"追加ページ {k + 1}：測定データの記録方法について。"])
        c.showPage()
    c.save()
    return path


def make_excel(path: Path, rows: int = 30) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import Reference, ScatterChart, Series
    from openpyxl.chart.trendline import Trendline

    wb = Workbook()
    ws = wb.active
    ws.title = "データ"
    ws.append(["荷重P(kN)", "変位(mm)", "応力度(N/mm2)", "断面積(mm2)"])
    ws["D2"] = 78.5
    for i in range(2, rows + 2):
        ws.cell(i, 1, (i - 1) * 2.0)
        ws.cell(i, 2, round((i - 1) * 0.025, 3))
        ws.cell(i, 3, f"=A{i}*1000/$D$2")
    ws.cell(rows + 3, 1, "平均")
    ws.cell(rows + 3, 3, f"=AVERAGE(C2:C{rows + 1})")
    ch = ScatterChart()
    ch.title = "荷重-変位曲線"
    ch.x_axis.title = "変位(mm)"
    ch.y_axis.title = "荷重(kN)"
    s = Series(Reference(ws, min_col=1, min_row=2, max_row=rows + 1),
               Reference(ws, min_col=2, min_row=2, max_row=rows + 1), title="試験体A")
    s.trendline = Trendline(trendlineType="linear", dispEq=True)
    ch.series.append(s)
    ws.add_chart(ch, "F2")
    ws2 = wb.create_sheet("説明")
    ws2["A1"] = "グラフは散布図で作成し、近似曲線を表示すること。"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def make_video(path: Path, workdir: Path, scenes: list[tuple[float, list[str], str]] | None = None,
               size: str = "640x360") -> Path:
    """静止画スライド（シーン）を連結した動画＋サイン波音声。scenes: [(秒数, 文字列行, 背景色)]"""
    scenes = scenes or [
        (15, ["STEP 1", "Record load and displacement"], "white"),
        (20, ["STEP 2", "=AVERAGE(B2:B10)"], "lightyellow"),
        (15, ["STEP 3", "Scatter chart"], "lightblue"),
    ]
    workdir.mkdir(parents=True, exist_ok=True)
    lst = workdir / "list.txt"
    w, h = map(int, size.split("x"))
    with open(lst, "w") as f:
        for i, (dur, lines, bg) in enumerate(scenes):
            img = text_image(workdir / f"scene{i}.png", lines, size=(w, h), bg=bg, font_size=36)
            f.write(f"file '{img}'\nduration {dur}\n")
        f.write(f"file '{img}'\n")
    total = sum(s[0] for s in scenes)
    path.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(["ffmpeg", "-nostdin", "-y", "-v", "error", "-f", "concat", "-safe", "0", "-i", str(lst),
                    "-f", "lavfi", "-i", f"sine=frequency=440:duration={total}",
                    "-vf", "fps=10,format=yuv420p", "-c:v", "libx264", "-preset", "ultrafast", "-c:a", "aac",
                    "-t", str(total), "-shortest", str(path)], check=True)
    return path


def make_zip(path: Path, members: dict[str, Path | bytes]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, src in members.items():
            zf.writestr(name, src if isinstance(src, bytes) else Path(src).read_bytes())
    return path


def make_zip_bomb(path: Path, mb: int = 60) -> Path:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        with zf.open("zeros.bin", "w") as f:
            chunk = b"\0" * (1024 * 1024)
            for _ in range(mb):
                f.write(chunk)
    return path


def make_all(root: Path) -> Path:
    """一式を root/source に作る。"""
    src = root / "source"
    tmp = root / "_tmp"
    make_pdf(src / "第1回/実験指導書.pdf")
    make_pdf(src / "第2回/補足資料.pdf", deadline="11月27日")
    text_image(src / "第1回/板書.png", ["σ = P / A", "ε = ΔL / L"])
    make_excel(src / "第1回/測定データ.xlsx")
    make_video(src / "第1回/講義動画.mp4", tmp / "video")
    inner = make_pdf(tmp / "inner.pdf")
    nested = make_zip(tmp / "nested.zip", {"deep/memo.txt": "ネストしたZIPの中のメモ。必ず確認すること。".encode()})
    make_zip(src / "配布資料.zip", {
        "資料/inner.pdf": inner,
        "資料/板書2.png": text_image(tmp / "b2.png", ["E = σ / ε"]),
        "資料/nested.zip": nested,
        "../evil.txt": b"path traversal",
        "/abs.txt": b"absolute",
        "~$lock.xlsx": b"temp",
    })
    return src


if __name__ == "__main__":
    import sys
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "dummy_data")
    print(make_all(out))
