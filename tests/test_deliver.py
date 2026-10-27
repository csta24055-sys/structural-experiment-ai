"""Excel・レポートの作成（RC 梁）。合成の入力（実際の実験データではない）で、式・数式・検査を確かめる。"""
import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent))
import rc_synth  # noqa: E402

from structlab.deliver import rc_beam as R  # noqa: E402

ROOT = Path(__file__).parent.parent
needs_lo = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice が無い")


def test_formulas_reproduce_the_printed_design_values():
    """書式 PDF の式で設計値（表1）から計算すると、印刷された FQ・剛性などと一致する（差は印刷の丸めだけ: 0.6% 以内）。"""
    t = R.theory(R.DESIGN)
    for k, v in R.PRINTED.items():
        assert abs(t[k] / v - 1) < 0.006, k
    assert t["mu"] < 1 and "せん断" in t["failure"]
    (x0, y0), (x1, y1), (x2, y2) = t["trilinear_mm_N"]
    assert abs(y1 / 1000 - 34.6) < 0.1 and abs(x1 - 0.28) < 0.01 and abs(y2 / 1000 - 76.2) < 0.2 and abs(x2 - 1.58) < 0.01


def test_inputs_are_read_and_gaps_are_reported_not_guessed(tmp_path):
    inp = R.load_inputs(rc_synth.write_inputs(tmp_path / "in"))
    s = R.check(inp)
    assert s["rows"] == 200 and inp.materials["sigma_B"] == 22.0 and inp.mapping["P"] == "荷重P"
    assert s["events"]["BC"]["matching_rows"] == [50]                   # データ.txt の変位 = 補正たわみ
    kinds = {i["what"].split(":")[-1].strip(): i["kind"] for i in inp.issues}
    assert any("せん断ひび割れ" in w and k == "不足" for w, k in kinds.items())    # 無い値は推測で補わない
    assert any("材料強度.txt" in i["what"] for i in inp.issues)


def test_corrected_deflection_mismatch_is_flagged(tmp_path):
    d = rc_synth.write_inputs(tmp_path / "in")
    t = (d / "データ.txt").read_bytes().decode("cp932").replace("変位 --> 0.98", "変位 --> 0.75")
    (d / "データ.txt").write_bytes(t.encode("cp932"))
    inp = R.load_inputs(d)
    R.check(inp)
    assert any(i["kind"] == "不明" and "補正たわみ" in i["what"] for i in inp.issues)


@needs_lo
def test_excel_formulas_match_independent_python(tmp_path):
    import openpyxl

    from structlab.deliver import excel as X
    inp = R.load_inputs(rc_synth.write_inputs(tmp_path / "in"))
    rc = X.recalc_with_libreoffice(X.build_workbook(inp, tmp_path / "src" / "a.xlsx"), tmp_path / "rc")
    wb = openpyxl.load_workbook(rc, data_only=True)
    exp = X.expected_values(inp)
    for row in wb["表4"].iter_rows(min_row=2, values_only=True):
        fq, cq = exp[row[0].split()[-1]]
        assert row[1] == pytest.approx(fq, rel=1e-9) and row[2] == pytest.approx(cq, rel=1e-9)
    d = wb["測定データ"]
    assert d["K2"].value == pytest.approx(inp.rows[0]["d1"] - (inp.rows[0]["d2"] + inp.rows[0]["d3"]) / 2)
    charts = openpyxl.load_workbook(rc)["図"]._charts
    assert len(charts) == 3
    assert str(wb["図"]["A20"].value).startswith("図-")                  # 図の題は図の下のセル


@needs_lo
def test_report_is_assembled_and_missing_items_block_completion(tmp_path):
    from structlab.deliver import report as RP
    inp = R.load_inputs(rc_synth.write_inputs(tmp_path / "in"))
    content = {"cover": {k: "x" for k in RP.COVER_FIELDS}, "table3_max": "Qmin",
               "sections": [{"heading": "結果", "blocks": [{"figure": "load_deflection"}, {"table": "表4"}]}],
               "discussion": {n: {k: "文" for k in it} for n, it in RP.MANDATORY.items()}}
    content["discussion"]["4"]["strain_graph"] = {"blocks": [{"figure": "strain_main"}, {"text": "文"}]}
    res = RP.build(inp, content, tmp_path / "out")
    assert res["pdf"].exists() and res["figures"] == 2
    assert any("表3" in i["what"] for i in res["issues"])                # 入力に無い荷重は【未記入】→ 完成扱いにしない
    del content["discussion"]["2"]["failure_mode"]
    content["cover"]["氏名"] = ""
    iss = RP.validate(content)
    assert any("考察 2)" in i["what"] for i in iss) and any("氏名" in i["what"] for i in iss)


def test_deliverables_need_explicit_approval(tmp_path):
    r = subprocess.run([sys.executable, str(ROOT / "scripts" / "deliver.py"), "excel", "--inputs", str(tmp_path)],
                       capture_output=True, text=True)
    assert r.returncode != 0 and "許可" in r.stderr
