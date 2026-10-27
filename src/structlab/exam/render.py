"""文書出力（PDF・Word）と品質検査。内容はブロックの列で受け取り、見た目はここだけで決める。

ブロック: {"t": "title"|"h1"|"h2"|"p"|"small"|"pre"|"list"|"table"|"space"|"break", ...}
  p/small/pre: text   list: items   table: rows（1 行目が見出し）   space: lines（解答欄の行数）
白背景・黒文字・標準的な日本語書体。日本語フォントは埋め込む（環境に無ければ CID フォント）。
"""
from __future__ import annotations

from pathlib import Path

FONT_CANDIDATES = [
    "/usr/share/fonts/opentype/ipaexfont-gothic/ipaexg.ttf",
    "/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf",
    "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
]
_FONT: str | None = None


def jp_font() -> str:
    global _FONT
    if _FONT:
        return _FONT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    for p in FONT_CANDIDATES:
        if Path(p).exists():
            try:
                pdfmetrics.registerFont(TTFont("JP", p, subfontIndex=0) if p.endswith(".ttc") else TTFont("JP", p))
                _FONT = "JP"
                return _FONT
            except Exception:  # noqa: BLE001
                continue
    pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    _FONT = "HeiseiKakuGo-W5"
    return _FONT


def _esc(s: str) -> str:
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>"))


def to_pdf(blocks: list[dict], path: Path, title: str, header: str = "") -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)
    f = jp_font()
    st = {
        "title": ParagraphStyle("title", fontName=f, fontSize=16, leading=22, alignment=TA_CENTER, spaceAfter=8),
        "h1": ParagraphStyle("h1", fontName=f, fontSize=13, leading=18, spaceBefore=8, spaceAfter=4),
        "h2": ParagraphStyle("h2", fontName=f, fontSize=11.5, leading=16, spaceBefore=6, spaceAfter=2),
        "p": ParagraphStyle("p", fontName=f, fontSize=10, leading=15, wordWrap="CJK"),
        "small": ParagraphStyle("small", fontName=f, fontSize=8.5, leading=12, wordWrap="CJK"),
        "pre": ParagraphStyle("pre", fontName=f, fontSize=9.5, leading=14, wordWrap="CJK", leftIndent=8),
    }
    story = []
    for b in blocks:
        t = b["t"]
        if t in st:
            story.append(Paragraph(_esc(b.get("text", "")), st[t]))
        elif t == "list":
            for it in b["items"]:
                story.append(Paragraph("・" + _esc(it), st["p"]))
        elif t == "table":
            rows = [[Paragraph(_esc(c), st["small"]) for c in r] for r in b["rows"]]
            tb = Table(rows, repeatRows=1, colWidths=b.get("widths"))
            tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.black),
                                    ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                                    ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(tb)
        elif t == "space":
            n = int(b.get("lines", 4))
            box = Table([[""]] * n, colWidths=[170 * mm], rowHeights=[8 * mm] * n)
            box.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.grey),
                                     ("BOX", (0, 0), (-1, -1), 0.6, colors.black)]))
            story.append(KeepTogether([box]))
        elif t == "break":
            story.append(PageBreak())
        story.append(Spacer(1, 2 if t in ("list", "p") else 4))

    def footer(c, d):
        c.saveState()
        c.setFont(f, 8)
        c.drawCentredString(A4[0] / 2, 10 * mm, f"- {d.page} -")
        if header:
            c.drawString(15 * mm, A4[1] - 10 * mm, header)
        c.restoreState()

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(path), pagesize=A4, title=title, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return path


def to_docx(blocks: list[dict], path: Path, title: str) -> Path:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "IPAGothic"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "IPAゴシック")
    style.font.size = Pt(10.5)
    d.core_properties.title = title
    for b in blocks:
        t = b["t"]
        if t == "title":
            p = d.add_heading(b["text"], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif t == "h1":
            d.add_heading(b["text"], level=1)
        elif t == "h2":
            d.add_heading(b["text"], level=2)
        elif t in ("p", "pre", "small"):
            p = d.add_paragraph(b["text"])
            if t == "small":
                for r in p.runs:
                    r.font.size = Pt(8.5)
        elif t == "list":
            for it in b["items"]:
                d.add_paragraph(it, style="List Bullet")
        elif t == "table":
            rows = b["rows"]
            tb = d.add_table(rows=len(rows), cols=len(rows[0]))
            tb.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    tb.cell(i, j).text = str(c)
        elif t == "space":
            for _ in range(int(b.get("lines", 4))):
                d.add_paragraph("_" * 60)
        elif t == "break":
            d.add_page_break()
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))
    return path


def check_pdf(path: Path, expect: list[str] = (), render_dir: Path | None = None) -> dict:
    """PDF の品質検査: 開けるか・各ページに文字があるか・文字化け（置換文字）・期待する文の有無・ページ番号。"""
    import pymupdf
    res = {"file": path.name, "ok": True, "problems": []}
    try:
        doc = pymupdf.open(str(path))
    except Exception as e:  # noqa: BLE001
        return {"file": path.name, "ok": False, "problems": [f"開けない: {e}"]}
    with doc:
        res["pages"] = doc.page_count
        texts = [p.get_text() for p in doc]
        fonts = {f[3] for p in doc for f in p.get_fonts()}
        if render_dir is not None and doc.page_count:
            render_dir.mkdir(parents=True, exist_ok=True)
            doc[0].get_pixmap(dpi=60).save(str(render_dir / (path.stem + "_p1.png")))
    allt = "".join(texts)
    if not res["pages"]:
        res["problems"].append("ページが無い")
    empty = [i + 1 for i, t in enumerate(texts) if not t.strip()]
    if empty:
        res["problems"].append(f"文字の無いページ {empty}")
    if "�" in allt:
        res["problems"].append("置換文字（文字化け）がある")
    flat = allt.replace("\n", "").replace(" ", "")
    missing = [e for e in expect if e.replace("\n", "").replace(" ", "") not in flat]
    if missing:
        res["problems"].append(f"期待する文が見つからない: {missing[:3]}")
    if texts and "- 1 -" not in texts[0]:
        res["problems"].append("ページ番号が無い")
    res["fonts"] = sorted(fonts)
    res["ok"] = not res["problems"]
    return res


def check_docx(path: Path, expect: list[str] = ()) -> dict:
    import docx
    try:
        d = docx.Document(str(path))
    except Exception as e:  # noqa: BLE001
        return {"file": path.name, "ok": False, "problems": [f"開けない: {e}"]}
    text = "\n".join(p.text for p in d.paragraphs) + "\n".join(
        c.text for t in d.tables for r in t.rows for c in r.cells)
    missing = [e for e in expect if e not in text]
    probs = [f"期待する文が見つからない: {missing[:3]}"] if missing else []
    if "�" in text:
        probs.append("置換文字（文字化け）がある")
    return {"file": path.name, "ok": not probs, "problems": probs}
