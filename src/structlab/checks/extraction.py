"""資料抽出時のチェック（既存の抽出結果を読み、不足していた検査だけを行う）。

- 音声と映像の時刻対応（発話が区間内か、対応付けた画面が発話時点で表示済みか）
- 文字起こし（音声があるのに発話ゼロ、低信頼区間、文字起こし未実施）
- OCR（文字化けの疑い、OCRなしのスキャン資料）
- 数値・数式（数値の誤読の疑い、括弧の対応）
- Excel（数式の再計算とキャッシュ値の照合、エラー値、空セル参照）
- 読み取り失敗（失敗した単位）
既存の検査（ダウンロード時のハッシュ照合、ZIPの安全性、スキャンページ判定など）は重複して行わない。
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from ..db import DB
from ..models import fmt_time
from .base import clear, record
from .formula import ERRORS, Evaluator, FormulaError, Unsupported, close_enough

STAGE = "extraction"
_OK_CHARS = re.compile(r"[぀-ヿ㐀-鿿０-ｚA-Za-z0-9\s.,:;()\[\]+\-*/=%、。・「」（）：＝σεΔπ°μ'\"!?！？]")
_RUN = re.compile(r"(.)\1{5,}")
# 数字と英字の O/o/l/I が混ざった数値らしい語（OCR・入力での誤読の兆候）
_NUM_SUSPECT = re.compile(r"(?<![A-Za-z0-9])(?=[0-9OolI.]*\d)(?=[0-9OolI.]*[OolI])[0-9OolI][0-9OolI.]*(?![A-Za-z0-9])")


def ocr_quality(text: str) -> tuple[str, str]:
    t = (text or "").strip()
    if not t:
        return "unverified", "文字を読み取れなかった"
    bad = sum(1 for ch in t if not _OK_CHARS.match(ch))
    ratio = bad / max(len(t), 1)
    run = _RUN.search(t)
    if ratio > 0.3 or run:
        why = f"想定外の記号の割合 {ratio:.0%}" + (f"／同じ文字の連続「{run.group(0)[:8]}」" if run else "")
        return "needs_review", f"文字化け・誤読の疑い（{why}）"
    return "ok", f"{len(t)} 文字"


def suspicious_numbers(text: str) -> list[str]:
    return sorted(set(_NUM_SUSPECT.findall(text or "")))


def unbalanced(s: str) -> bool:
    depth = 0
    for ch in s:
        if ch in "(（":
            depth += 1
        elif ch in ")）":
            depth -= 1
            if depth < 0:
                return True
    return depth != 0


def check_file(db: DB, cfg, frow, local_path: Path | None, run_id: int | None = None) -> None:
    fid = frow["id"]
    clear(db, STAGE, file_id=fid)
    units = db.q("SELECT * FROM units WHERE file_id=? AND file_hash=? ORDER BY seq, id", (fid, frow["content_hash"] or ""))
    with db.tx():
        for u in units:
            lo = json.loads(u["locator"])
            if u["status"] == "failed":
                record(db, stage=STAGE, check_type="read", target_type="unit", target_id=u["id"], status="failed",
                       detail=f"読み取りに失敗: {u['last_error']}", file_id=fid, unit_id=u["id"], run_id=run_id)
                continue
            rec = json.loads(u["extract_json"] or "{}")
            data = rec.get("data") or {}
            kw = dict(file_id=fid, unit_id=u["id"], run_id=run_id)
            if u["unit_type"] == "video_segment":
                _video(db, u, lo, data, kw)
            elif u["unit_type"] == "pdf_page":
                _pdf(db, u, lo, rec, data, kw)
            elif u["unit_type"] == "image":
                st, d = ocr_quality(rec.get("text", "")) if data.get("ocr", "なし") != "なし" else ("unverified", "OCRなし（AIでの読み取りが必要）")
                record(db, stage=STAGE, check_type="ocr", target_type="unit", target_id=u["id"], status=st, detail=d, **kw)
    if frow["kind"] == "excel" and local_path is not None and Path(local_path).exists():
        recalc_workbook(db, cfg, frow, Path(local_path), run_id)


def _video(db, u, lo, data, kw) -> None:
    s, e = float(lo["start"]), float(lo["end"])
    rows = db.q("SELECT * FROM timeline WHERE unit_id=? ORDER BY t_start", (u["id"],))
    where = f"{fmt_time(s)}–{fmt_time(e)}"
    if lo.get("has_audio"):
        speech = [r for r in rows if r["speech"]]
        bad = []
        for r in speech:
            if not (s - 0.5 <= r["t_start"] <= e + 0.5):
                bad.append(f"発話 {fmt_time(r['t_start'])} が区間外")
            if r["frame_time"] is not None and r["frame_time"] > r["t_start"] + 0.5:
                bad.append(f"発話 {fmt_time(r['t_start'])} に、その後の画面 {fmt_time(r['frame_time'])} を対応付けている")
            if r["frame_time"] is not None and not (s - 0.5 <= r["frame_time"] <= e + 0.5):
                bad.append(f"画面 {fmt_time(r['frame_time'])} が区間外")
        if speech and lo.get("has_video"):
            record(db, stage=STAGE, check_type="av_sync", target_type="unit", target_id=u["id"],
                   status="mismatch" if bad else "ok", expected="発話開始時点で表示されていた画面を対応付ける",
                   actual=bad or f"{len(speech)} 発話を確認", detail=where, **kw)
        transcriber = data.get("transcriber", "none")
        if transcriber == "none":
            st, d = "unverified", "文字起こし未実施（transcription.backend=none）"
        elif not speech:
            st, d = "needs_review", "音声のある区間だが発話を検出できなかった（無音、または文字起こしの失敗の可能性）"
        elif data.get("low_confidence_speech"):
            st, d = "needs_review", "文字起こしの信頼度が低い発話: " + "、".join(fmt_time(t) for t in data["low_confidence_speech"][:20])
        else:
            st, d = "ok", f"{len(speech)} 発話（{transcriber}）"
        record(db, stage=STAGE, check_type="transcription", target_type="unit", target_id=u["id"], status=st,
               detail=f"{where} {d}", **kw)
    if lo.get("has_video"):
        n = data.get("n_keyframes", 0)
        record(db, stage=STAGE, check_type="frames", target_type="unit", target_id=u["id"],
               status="ok" if n else "needs_review", detail=f"{where} キーフレーム {n} 枚", **kw)
        ocrs = [r["screen_text"] for r in rows if r["screen_text"]]
        if data.get("ocr", "なし") == "なし":
            record(db, stage=STAGE, check_type="ocr", target_type="unit", target_id=u["id"], status="unverified",
                   detail=f"{where} 画面のOCRなし（画面の文字はAIでのみ読み取る）", **kw)
        elif ocrs:
            worst = max((ocr_quality(t) for t in ocrs), key=lambda x: {"ok": 0, "unverified": 1, "needs_review": 2}[x[0]])
            record(db, stage=STAGE, check_type="ocr", target_type="unit", target_id=u["id"], status=worst[0],
                   detail=f"{where} {worst[1]}", **kw)


def _pdf(db, u, lo, rec, data, kw) -> None:
    p = lo.get("page")
    text = rec.get("text", "")
    if data.get("chars", 0) < 30 and (data.get("n_images") or (data.get("n_drawings") or 0) > 0):
        if data.get("ocr_text") is not None:
            st, d = ocr_quality(data.get("ocr_text", ""))
            d = "スキャンページのOCR: " + d
        else:
            st, d = "unverified", "スキャンページでOCRなし（本文はAIでのみ読み取る）"
        record(db, stage=STAGE, check_type="ocr", target_type="unit", target_id=u["id"], status=st, detail=f"p.{p} {d}", **kw)
    sus = suspicious_numbers(text)
    record(db, stage=STAGE, check_type="numbers", target_type="unit", target_id=u["id"],
           status="needs_review" if sus else "ok",
           detail=f"p.{p} " + (f"数値の誤読の疑い（英字のO/l/Iが混入）: {', '.join(sus[:10])}" if sus else "数値表記に誤読の兆候なし"), **kw)
    bad = [f for f in data.get("formula_candidates", []) if unbalanced(f)]
    if data.get("formula_candidates"):
        record(db, stage=STAGE, check_type="formula", target_type="unit", target_id=u["id"],
               status="needs_review" if bad else "ok",
               detail=f"p.{p} " + (f"括弧の対応が取れない式: {' / '.join(bad[:5])}" if bad else
                                  f"数式候補 {len(data['formula_candidates'])} 行（構文の目視確認は未実施）"), **kw)


# ---------------- Excel の再計算 ----------------
def load_workbook_cells(path: Path, max_cells: int) -> tuple[dict, dict, bool]:
    """(数式 {(sheet,cell): '=...'}, 値 {(sheet,cell): value}, 上限で打ち切ったか)"""
    import openpyxl
    formulas, values = {}, {}
    truncated = False
    wf = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_vba=False)
    wv = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_vba=False)
    try:
        n = 0
        for ws in wf.worksheets:
            vs = wv[ws.title]
            for rf, rv in zip(ws.iter_rows(), vs.iter_rows(values_only=True)):
                for cf, v in zip(rf, rv):
                    raw = getattr(cf, "value", None)
                    coord = getattr(cf, "coordinate", None)
                    if coord is None or (raw is None and v is None):
                        continue
                    n += 1
                    if n > max_cells:
                        truncated = True
                        break
                    key = (ws.title, coord)
                    if isinstance(raw, str) and raw.startswith("="):
                        formulas[key] = raw
                        values[key] = v  # Excel が保存した計算結果（無ければ None）
                    else:
                        values[key] = raw if raw is not None else v
                if truncated:
                    break
            if truncated:
                break
    finally:
        wf.close()
        wv.close()
    return formulas, values, truncated


def _resolver(values: dict, formulas: dict | None = None, memo: dict | None = None, stack: set | None = None):
    """参照先が数式セルなら、Excel の保存値を使わず入力値から再帰的に計算し直す（独立した再計算）。"""
    formulas = formulas or {}
    memo = memo if memo is not None else {}
    stack = stack if stack is not None else set()

    def resolve(s: str, c: str):
        key = (s, c)
        if key in formulas:
            if key in memo:
                v = memo[key]
                if isinstance(v, Exception):
                    raise v
                return v
            if key in stack:
                raise Unsupported(f"循環参照 {s}!{c}")
            stack.add(key)
            try:
                v = Evaluator(resolve, s).evaluate(formulas[key])
                memo[key] = v
                return v
            except (FormulaError, Unsupported) as e:
                memo[key] = e
                raise
            finally:
                stack.discard(key)
        v = values.get(key)
        if hasattr(v, "isoformat"):
            raise Unsupported("日付・時刻のセルを参照")
        return v
    return resolve


def recalc_workbook(db: DB, cfg, frow, path: Path, run_id: int | None = None, *, stage: str = STAGE,
                    target_type: str | None = None, prefix: str | None = None) -> dict:
    """全数式を保存値から独立に再計算し、Excel が保存した計算結果と照合する。"""
    fid = frow["id"] if frow["id"] >= 0 else None
    ttype = target_type or "file"
    prefix = prefix if prefix is not None else str(fid)
    max_cells = int(cfg.data.get("checks", {}).get("max_excel_cells", 300000))
    max_rows = int(cfg.data.get("checks", {}).get("max_detail_rows", 500))
    try:
        formulas, values, truncated = load_workbook_cells(path, max_cells)
    except Exception as e:  # noqa: BLE001
        record(db, stage=stage, check_type="excel_recalc", target_type=ttype, target_id=prefix, status="failed",
               detail=f"ブックを読めない: {e}", file_id=fid, run_id=run_id)
        return {}
    counts = {"ok": 0, "mismatch": 0, "needs_review": 0, "unverified": 0}
    detail_n = 0
    memo: dict = {}
    with db.tx():
        for (sh, cell), f in sorted(formulas.items()):
            cached = values.get((sh, cell))
            ev = Evaluator(_resolver(values, formulas, memo, set()), sh)
            target = f"{sh}!{cell}"
            try:
                got = ev.evaluate(f)
                empty = [f"{s}!{c}" for s, c in ev.refs if values.get((s, c)) is None and (s, c) not in formulas]
                if cached is None:
                    st, d = "unverified", "Excelで計算された値が保存されていない（Excelで開いて保存すると照合できる）"
                elif isinstance(cached, str) and cached in ERRORS:
                    st, d = "mismatch", f"保存値がエラー {cached}"
                elif close_enough(got, cached):
                    st, d = ("needs_review", f"空のセルを参照: {', '.join(empty[:5])}") if empty and len(ev.refs) <= 3 else ("ok", "")
                else:
                    st, d = "mismatch", "再計算の結果が保存値と異なる"
            except FormulaError as e:
                got = str(e)
                st = "mismatch" if not (isinstance(cached, str) and cached == str(e)) else "needs_review"
                d = f"再計算でエラー {e}" + ("（保存値も同じエラー）" if st == "needs_review" else "")
            except Unsupported as e:
                got, st, d = None, "unverified", f"再計算に未対応: {e}"
            counts[st] += 1
            if st != "ok" and detail_n < max_rows:
                detail_n += 1
                record(db, stage=stage, check_type="excel_cell", target_type=target_type or "cell", target_id=f"{prefix}:{target}",
                       status=st, expected=cached, actual=got, detail=f"{target} {f} — {d}", file_id=fid, run_id=run_id)
        overall = ("mismatch" if counts["mismatch"] else "needs_review" if counts["needs_review"]
                   else "unverified" if counts["unverified"] or truncated or not formulas else "ok")
        record(db, stage=stage, check_type="excel_recalc", target_type=ttype, target_id=prefix, status=overall,
               expected="Excelの保存値", actual=counts,
               detail=f"数式 {len(formulas)} 個を再計算: 一致 {counts['ok']}・不一致 {counts['mismatch']}・要確認 "
                      f"{counts['needs_review']}・未検証 {counts['unverified']}" + ("／セル数の上限で打ち切り" if truncated else ""),
               file_id=fid, run_id=run_id)
    return counts
