"""Excel (.xlsx/.xlsm): シート構造・セル・数式・グラフ定義を調査する。マクロは実行しない。"""
from __future__ import annotations

import posixpath
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from ..models import UnitSpec
from .base import Ctx, Extraction, Processor

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
MAX_XML = 50 * 1024 * 1024
CHART_TYPES = {
    "lineChart": "折れ線", "scatterChart": "散布図", "barChart": "棒", "bar3DChart": "3D棒",
    "pieChart": "円", "areaChart": "面", "radarChart": "レーダー", "bubbleChart": "バブル",
    "doughnutChart": "ドーナツ", "line3DChart": "3D折れ線", "stockChart": "株価",
}


def _read_xml(zf: zipfile.ZipFile, name: str):
    info = zf.getinfo(name)
    if info.file_size > MAX_XML:
        raise ValueError(f"XMLが大きすぎます: {name}")
    return ET.fromstring(zf.read(name))


def _rels(zf: zipfile.ZipFile, part: str) -> dict[str, str]:
    d, base = posixpath.split(part)
    rel_name = posixpath.join(d, "_rels", base + ".rels")
    if rel_name not in zf.namelist():
        return {}
    root = _read_xml(zf, rel_name)
    out = {}
    for r in root.findall("rel:Relationship", NS):
        tgt = r.get("Target", "")
        out[r.get("Id")] = tgt.lstrip("/") if tgt.startswith("/") else posixpath.normpath(posixpath.join(d, tgt))
    return out


def _texts(el) -> str:
    return "".join(t.text or "" for t in el.iter(f"{{{NS['a']}}}t")) if el is not None else ""


def parse_chart(zf: zipfile.ZipFile, part: str) -> dict:
    root = _read_xml(zf, part)
    chart = root.find("c:chart", NS)
    info: dict = {"part": part, "title": _texts(chart.find("c:title", NS)) if chart is not None else "",
                  "types": [], "series": [], "axes": []}
    pa = chart.find("c:plotArea", NS) if chart is not None else None
    if pa is None:
        return info
    for child in pa:
        tag = child.tag.split("}")[1]
        if tag in CHART_TYPES:
            info["types"].append(CHART_TYPES[tag])
            for ser in child.findall("c:ser", NS):
                s = {}
                for key in ("tx", "cat", "val", "xVal", "yVal"):
                    el = ser.find(f"c:{key}", NS)
                    if el is not None:
                        f = el.find(".//c:f", NS)
                        v = el.find(".//c:v", NS)
                        s[key] = f.text if f is not None else (v.text if v is not None else "")
                tl = ser.find("c:trendline/c:trendlineType", NS)
                if tl is not None:
                    s["trendline"] = tl.get("val")
                    disp = ser.find("c:trendline/c:dispEq", NS)
                    s["trendline_show_equation"] = disp is not None and disp.get("val") in ("1", "true")
                info["series"].append(s)
        elif tag in ("valAx", "catAx", "dateAx", "serAx"):
            info["axes"].append({"type": tag, "title": _texts(child.find("c:title", NS))})
    return info


def workbook_structure(path: Path) -> dict:
    out: dict = {"sheets": [], "charts": [], "defined_names": [], "has_macro": False}
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        out["has_macro"] = "xl/vbaProject.bin" in names
        wb = _read_xml(zf, "xl/workbook.xml")
        wb_rels = _rels(zf, "xl/workbook.xml")
        for s in wb.findall("m:sheets/m:sheet", NS):
            rid = s.get(f"{{{NS['r']}}}id")
            out["sheets"].append({"name": s.get("name"), "state": s.get("state", "visible"),
                                  "part": wb_rels.get(rid, "")})
        for dn in wb.findall("m:definedNames/m:definedName", NS):
            out["defined_names"].append({"name": dn.get("name"), "ref": dn.text})
        for sh in out["sheets"]:
            if not sh["part"] or sh["part"] not in names:
                continue
            for _, drawing in _rels(zf, sh["part"]).items():
                if "drawings/" not in drawing or drawing not in names:
                    continue
                for _, ch in _rels(zf, drawing).items():
                    if "charts/chart" in ch and ch in names:
                        c = parse_chart(zf, ch)
                        c["sheet"] = sh["name"]
                        out["charts"].append(c)
            if sh["part"].startswith("xl/chartsheets/"):
                sh["chartsheet"] = True
    return out


def _norm_formula(f: str) -> str:
    """コピーされた数式をまとめるため、行番号を # に置換。"""
    return re.sub(r"(\$?[A-Z]{1,3})\$?\d+", r"\1#", f)


class ExcelProcessor(Processor):
    kind = "excel"

    def plan(self, path: Path, ctx: Ctx) -> list[UnitSpec]:
        import openpyxl

        rows_per = int(ctx.cfg.data["excel"]["rows_per_unit"])
        units = [UnitSpec("excel_workbook", {"part": "structure"}, 0)]
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_vba=False)
        try:
            seq = 1
            for ws in wb.worksheets:
                max_row = ws.max_row
                if max_row is None:  # 寸法情報が無いファイルは走査して数える（ストリーミング）
                    max_row = sum(1 for _ in ws.iter_rows(values_only=True))
                max_row = max(max_row or 1, 1)
                for start in range(1, max_row + 1, rows_per):
                    units.append(UnitSpec("excel_sheet", {"sheet": ws.title, "row_start": start,
                                                          "row_end": min(start + rows_per - 1, max_row)}, seq))
                    seq += 1
        finally:
            wb.close()
        return units

    def label(self, unit: UnitSpec) -> str:
        if unit.unit_type == "excel_workbook":
            return "ブック構造"
        lo = unit.locator
        return f"シート「{lo['sheet']}」{lo['row_start']}〜{lo['row_end']}行"

    def extract(self, path: Path, unit: UnitSpec, ctx: Ctx) -> Extraction:
        if unit.unit_type == "excel_workbook":
            return self._structure(path)
        return self._sheet(path, unit, ctx)

    def _structure(self, path: Path) -> Extraction:
        st = workbook_structure(path)
        ex = Extraction(data=st)
        lines = ["### シート", *[f"- {s['name']}（{s['state']}）" for s in st["sheets"]]]
        if st["defined_names"]:
            lines += ["### 名前定義", *[f"- {d['name']} = {d['ref']}" for d in st["defined_names"]]]
        lines.append("### グラフ")
        for c in st["charts"]:
            lines.append(f"- シート「{c.get('sheet')}」: 種類={'/'.join(c['types']) or '不明'} タイトル={c['title'] or '(なし)'}")
            for ax in c["axes"]:
                lines.append(f"  - 軸 {ax['type']}: {ax['title'] or '(タイトルなし)'}")
            for s in c["series"]:
                lines.append(f"  - 系列: {s}")
        if not st["charts"]:
            lines.append("（グラフなし）")
        if st["has_macro"]:
            lines.append("### 注意\nマクロ(VBA)を含みます。本ツールはマクロを実行しません。")
        ex.ai_text = "\n".join(lines)
        ex.text = ex.ai_text
        ex.segments = [{"text": "\n".join(c["title"] for c in st["charts"] if c["title"]), "origin": "document_text"}]
        return ex

    def _sheet(self, path: Path, unit: UnitSpec, ctx: Ctx) -> Extraction:
        import openpyxl

        lo = unit.locator
        max_cells = int(ctx.cfg.data["excel"]["max_cells_per_unit"])
        wf = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_vba=False)
        wv = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_vba=False)
        cells: list[dict] = []
        formulas: dict[str, dict] = {}
        truncated = False
        try:
            fs, vs = wf[lo["sheet"]], wv[lo["sheet"]]
            it_f = fs.iter_rows(min_row=lo["row_start"], max_row=lo["row_end"])
            it_v = vs.iter_rows(min_row=lo["row_start"], max_row=lo["row_end"], values_only=True)
            for row_f, row_v in zip(it_f, it_v):
                for cf, v in zip(row_f, row_v):
                    raw = getattr(cf, "value", None)
                    if raw is None:
                        continue
                    coord = getattr(cf, "coordinate", None)
                    if coord is None:
                        continue
                    entry = {"cell": coord}
                    if isinstance(raw, str) and raw.startswith("="):
                        entry["formula"] = raw
                        entry["value"] = v
                        key = _norm_formula(raw)
                        g = formulas.setdefault(key, {"pattern": key, "example": raw, "example_cell": coord,
                                                      "count": 0})
                        g["count"] += 1
                    else:
                        entry["value"] = raw if not hasattr(raw, "isoformat") else raw.isoformat()
                    if len(cells) < max_cells:
                        cells.append(entry)
                    else:
                        truncated = True
        finally:
            wf.close()
            wv.close()
        ex = Extraction(data={"sheet": lo["sheet"], "row_start": lo["row_start"], "row_end": lo["row_end"],
                              "n_cells": len(cells), "truncated": truncated,
                              "formula_groups": list(formulas.values())})
        if truncated:
            ex.needs_ai.append(f"セル数が上限({max_cells})を超えたため一部のみ記録（excel.max_cells_per_unit）")
        text_cells = [c for c in cells if isinstance(c.get("value"), str) and "formula" not in c]
        ex.text = "\n".join(str(c["value"]) for c in text_cells) + "\n" + "\n".join(
            c["formula"] for c in cells if "formula" in c)
        ex.segments = [{"text": "\n".join(str(c["value"]) for c in text_cells), "origin": "document_text"}]
        preview = cells[:400]
        lines = [f"### シート「{lo['sheet']}」{lo['row_start']}〜{lo['row_end']}行（非空セル {len(cells)}）",
                 "### 数式パターン（同型の数式をまとめたもの）"]
        lines += [f"- {g['example_cell']}: {g['example']} （同型 {g['count']} セル）" for g in formulas.values()] or ["（なし）"]
        lines.append("### セル（先頭400件）")
        for c in preview:
            lines.append(f"{c['cell']}: " + (f"{c['formula']} → {c.get('value')}" if "formula" in c else repr(c["value"])))
        ex.ai_text = "\n".join(lines)
        return ex
