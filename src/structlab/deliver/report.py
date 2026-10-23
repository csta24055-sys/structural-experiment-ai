"""レポートの組み立て（RC 梁の破壊実験）: 図（matplotlib）・表（計算値）・本文（content.json）→ docx → PDF（LibreOffice）。

本文（考察など）はこのモジュールでは書かない。本番では、G5 の解析で得た講義の指示と計算結果をもとに Claude Code が
content.json を書き、このモジュールが書式どおりに組み立てて検査する。必須（◎）の項目が欠けている・「未記入」の印が残る・
表紙の欄が空のときは issues に残し、完成扱いにしない。
講義の指示: 図の題は図の下（「グラフとか図のタイトルは下って決まってる」）、散布図、凡例。表の題は書式 p9 の例どおり表の上。
"""
from __future__ import annotations

import json
from pathlib import Path

from .rc_beam import DESIGN, Inputs, corrected, measured_params, theory

#: 書式 p7（Ⅴ 考察）の必須（◎）の項目。content.json の discussion[番号][キー] に本文を書く
MANDATORY = {
    "1": {"correction": "たわみ測定値の補正方法（なぜ補正が必要か・採用する補正方法）",
          "relation": "補正したたわみ測定値と理論値の関係性（理論値はどうあるべきか）",
          "failure_mode": "梁の破壊モード（理論値のグラフとの比較から言える範囲で）"},
    "2": {"relation": "各種実強度と理論値の関係性（両者の比の大小に着目。理論値はどうあるべきか）",
          "failure_mode": "梁の破壊モード（最大耐力と矛盾しないように）"},
    "3": {"relation": "各種計算値と理論値の関係性（計算式に対する材料強度の影響度合いに着目）"},
    "4": {"strain_graph": "ひずみ度の測定結果のグラフと梁の破壊状況（グラフの傾きに着目）",
          "failure_mode": "梁の破壊モード（ひずみ度のグラフと矛盾しないように）"},
}
COVER_FIELDS = ("実験日時", "班", "学生番号", "氏名", "講義日", "実験日", "提出日")
PLACEHOLDER = "【未記入】"


def _font():
    import matplotlib
    from matplotlib import font_manager
    for name in ("Noto Sans CJK JP", "IPAexGothic", "IPAGothic", "WenQuanYi Zen Hei", "Unifont-JP"):
        if any(f.name == name for f in font_manager.fontManager.ttflist):
            matplotlib.rcParams["font.family"] = name
            return name
    return None


#: 鉄筋のヤング係数（材料試験の結果に無いので一般値 2.05×10⁵ N/mm² を使う。降伏ひずみの目安 εy = σy/Es の計算だけに使う）
ES_ASSUMED = 2.05e5


def _jp_font_file() -> str | None:
    import subprocess
    for name in ("IPAGothic", "Noto Sans CJK JP", "IPAexGothic"):
        try:
            f = subprocess.run(["fc-match", "-f", "%{file}", name], capture_output=True, text=True, timeout=10).stdout.strip()
        except (OSError, subprocess.SubprocessError):
            return None
        if f:
            return f
    return None


def figures(inp: Inputs, outdir: Path) -> dict[str, dict]:
    """図を PNG で書く。題は図の下に置くため画像には入れず、キャプションとして返す。
    荷重–補正たわみは書式 p8 の図－2 の描き方（理論値の三折れ線・2×CQ の水平線・上に部材角 R・右に τ）に合わせる。"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from .rc_beam import valid_deflection_rows
    _font()
    outdir.mkdir(parents=True, exist_ok=True)
    c = theory(measured_params(inp.materials))
    f = theory(DESIGN)
    a, b, j = DESIGN["a"], DESIGN["b"], DESIGN["j"]
    nv = valid_deflection_rows(inp.rows)
    rows_v = inp.rows[:nv]
    P = [r["P"] for r in rows_v]
    dl = [corrected(r) for r in rows_v]
    out = {}
    dy_c, py_c = 2 * c["Qy"] / c["Ky"], 2 * c["Qy"] / 1000

    def load_deflection(xmax: float, name: str, cap: str):
        fig, ax = plt.subplots(figsize=(7.2, 5.0), dpi=200)
        ax.plot(dl, P, "-", lw=0.6, color="tab:blue", alpha=0.6)
        ax.scatter(dl, P, s=3, color="tab:blue", label="実験値（補正たわみ δ）", zorder=3)
        xs, ys = zip(*c["trilinear_mm_N"])
        ax.plot(xs, [y / 1000 for y in ys], "-", lw=2.4, color="black", label="理論値 CQ（三折れ線）", zorder=4)
        ax.scatter(xs[1:], [y / 1000 for y in ys[1:]], s=25, color="black", zorder=5)
        if ys[-1] / 1000 < py_c - 1e-9:   # せん断破壊点から曲げ降伏点までは点線（書式の図）
            ax.plot([xs[-1], dy_c], [ys[-1] / 1000, py_c], ":", lw=2, color="black")
            ax.scatter([dy_c], [py_c], s=25, color="black", zorder=5)
            # せん断破壊後の耐力低下（書式 p8 の図－2 の下向きの破線矢印）
            ax.annotate("", xy=(xs[-1] + 0.35 * dy_c, ys[-1] / 1000 * 0.8), xytext=(xs[-1], ys[-1] / 1000),
                        arrowprops=dict(arrowstyle="->", ls="--", color="black", lw=1))
        ax.plot([0, dy_c * 1.25], [0, py_c * 1.25], "-", lw=0.7, color="gray", label="降伏時の剛性 Ky（理論値）")
        fx, fy = zip(*f["trilinear_mm_N"])
        ax.plot(fx, [y / 1000 for y in fy], "--", lw=1.2, color="tab:orange", label="計算値 FQ（三折れ線）")
        for sym, lab in (("Qy", "2×CQy"), ("Qmean", "2×CQmean"), ("Qmin", "2×CQmin"), ("QSC", "2×CQSC"), ("QBC", "2×CQBC")):
            v = 2 * c[sym] / 1000
            ax.axhline(v, lw=0.5, color="gray", ls="-")
            ax.text(xmax * 0.995, v + 0.6, f"{lab} = {v:.1f} kN", ha="right", va="bottom", fontsize=6.5)
        for key, mk, lab in (("BC", "^", "曲げひび割れ"), ("SC", "s", "せん断ひび割れ"), ("y", "D", "曲げ降伏")):
            e = inp.events.get(key)
            if e and e.get("delta_mm") is not None:
                ax.scatter([e["delta_mm"]], [e["P_kN"]], marker=mk, s=45, color="tab:red", edgecolor="black",
                           label=f"{lab}（データ.txt）", zorder=6)
        if "max" in inp.events and inp.rows:
            im = max(range(len(inp.rows)), key=lambda i: inp.rows[i]["P"])
            ax.scatter([corrected(inp.rows[im])], [inp.rows[im]["P"]], marker="*", s=90, color="gold", edgecolor="black",
                       label="最大荷重（測定データの最大の行）", zorder=6)
        ax.set_xlim(0, xmax)
        ax.set_ylim(0, 115)
        ax.set_xlabel("補正たわみ δ（mm）")
        ax.set_ylabel("荷重 P（kN）")
        sx = ax.secondary_xaxis("top", functions=(lambda x: x / a, lambda r: r * a))
        ticks = [t for t in (1 / 500, 1 / 250, 1 / 200, 1 / 150, 1 / 100, 1 / 75, 1 / 50) if t * a <= xmax]
        sx.set_xticks(ticks)
        sx.set_xticklabels([f"1/{round(1 / t)}" for t in ticks], fontsize=7)
        sx.set_xlabel("部材角 R（rad）")
        sy = ax.secondary_yaxis("right", functions=(lambda p: p * 1000 / (2 * b * j), lambda t: t * 2 * b * j / 1000))
        sy.set_ylabel("平均せん断応力度 τ（N/mm²）")
        ax.legend(fontsize=7, loc="upper center", bbox_to_anchor=(0.5, -0.12), ncol=3, frameon=False)
        ax.grid(alpha=0.25)
        fig.savefig(outdir / f"fig_{name}.png", bbox_inches="tight")
        plt.close(fig)
        out[name] = {"path": outdir / f"fig_{name}.png", "caption": cap}

    note = "" if nv == len(inp.rows) else f"（変位計 No.3 の計測が途切れる測定番号 {nv + 1} より前のデータ）"
    xfull = max(3.0, (int(max(dl) if dl else 3) + 1))
    load_deflection(xfull, "load_deflection", f"荷重 P と補正たわみ δ の関係（実験値と理論値・計算値）{note}")
    load_deflection(3.0, "load_deflection_zoom", "荷重 P と補正たわみ δ の関係（δ ≦ 3 mm の拡大。書式の図－2 と同じ範囲）")

    Pall = [r["P"] for r in inp.rows]
    lines = [(inp.events[k]["P_kN"], lab) for k, lab in (("y", "曲げ降伏（データ.txt）"), ("max", "最大荷重")) if k in inp.events]
    for key, cols, sig, cap in (("strain_main", (("e_main1", "主筋 No.1"),), inp.materials.get("main_sigma_y"),
                                 "荷重 P と主筋のひずみ度の関係"),
                                ("strain_stirrup", (("e_st2", "あばら筋 No.2"), ("e_st3", "あばら筋 No.3"), ("e_st4", "あばら筋 No.4")),
                                 inp.materials.get("stirrup_sigma_y"), "荷重 P とあばら筋のひずみ度の関係")):
        fig, ax = plt.subplots(figsize=(7.2, 4.6), dpi=200)
        for col, name in cols:
            ax.plot([r[col] for r in inp.rows], Pall, "-", lw=0.5, alpha=0.5)
            ax.scatter([r[col] for r in inp.rows], Pall, s=3, label=name)
        if sig:
            ey = sig / ES_ASSUMED * 1e6
            ax.axvline(ey, color="red", ls="--", lw=1, label=f"降伏ひずみの目安 σy/Es = {ey:,.0f}μ（Es = 2.05×10$^5$ N/mm² と仮定）")
        for n, (v, lab) in enumerate(lines):
            ax.axhline(v, color="gray", lw=0.6, ls=":")
            ax.text(ax.get_xlim()[1] * 0.99, v + (0.8 if n else -0.8), f"{lab} {v:.2f} kN", ha="right",
                    va="bottom" if n else "top", fontsize=6.5)
        ax.set_xlim(left=min(0, ax.get_xlim()[0]))
        ax.set_ylim(0, 115)
        ax.set_xlabel("ひずみ度 ε（μ）")
        ax.set_ylabel("荷重 P（kN）")
        ax.legend(fontsize=7, loc="lower right", framealpha=0.95)
        ax.grid(alpha=0.25)
        fig.tight_layout()
        fig.savefig(outdir / f"fig_{key}.png")
        plt.close(fig)
        out[key] = {"path": outdir / f"fig_{key}.png", "caption": cap}
    return out


def tables(inp: Inputs) -> dict[str, dict]:
    c, f = theory(measured_params(inp.materials)), theory(DESIGN)
    m, ev = inp.materials, inp.events
    t2 = {"title": "コンクリートおよび鋼の試験結果", "head": ["材料", "項目", "使用材料の実強度 [N/mm²]", "設計時強度 [N/mm²]"],
          "rows": [["コンクリート", "圧縮強度 σB", f"{m['sigma_B']:.1f}", "21"], ["主筋", "降伏点 σy", f"{m['main_sigma_y']:.1f}", "380"],
                   ["主筋", "引張強さ σm", f"{m.get('main_sigma_m', float('nan')):.1f}", "—"],
                   ["あばら筋", "降伏点 σwy", f"{m['stirrup_sigma_y']:.1f}", "325"],
                   ["あばら筋", "引張強さ σwm", f"{m.get('stirrup_sigma_m', float('nan')):.1f}", "—"]]}
    rows3 = []
    for label, key, sym in (("曲げひび割れ発生時せん断力 QBC", "BC", "QBC"), ("せん断ひび割れ発生時せん断力 QSC", "SC", "QSC"),
                            ("曲げ降伏時耐力 Qy", "y", "Qy")):
        tq = ev[key]["P_kN"] / 2 if key in ev else None
        rows3.append([label, f"{tq:.2f}" if tq is not None else PLACEHOLDER, f"{c[sym] / 1000:.2f}",
                      f"{tq / (c[sym] / 1000):.2f}" if tq is not None else PLACEHOLDER])
    t3 = {"title": "実強度（実験結果）と理論値の比較", "head": ["", "実強度 TQ [kN]", "理論値 CQ [kN]", "TQ/CQ"], "rows": rows3,
          "note": "最大耐力の行は、比べる理論値（書式に指定なし）を決めて content.json の table3_max で指定する"}
    t4 = {"title": "計算値と理論値の比較", "head": ["", "設計強度による計算値 FQ [kN]", "実材料強度による理論値 CQ [kN]", "FQ/CQ"],
          "rows": [[label, f"{f[s] / 1000:.2f}", f"{c[s] / 1000:.2f}", f"{f[s] / c[s]:.2f}"]
                   for label, s in (("曲げひび割れ発生時せん断力 QBC", "QBC"), ("せん断ひび割れ発生時せん断力 QSC", "QSC"),
                                    ("曲げ降伏時耐力 Qy", "Qy"), ("せん断耐力 min 式 Qmin", "Qmin"),
                                    ("せん断耐力 mean 式 Qmean", "Qmean"))]}
    return {"表2": t2, "表3": t3, "表4": t4}


def validate(content: dict) -> list[dict]:
    issues = []
    cover = content.get("cover", {})
    for k in COVER_FIELDS:
        if not str(cover.get(k, "")).strip():
            issues.append({"kind": "不足", "what": f"表紙の「{k}」が空"})
    disc = content.get("discussion", {})
    for no, items in MANDATORY.items():
        for key, what in items.items():
            text = str(disc.get(no, {}).get(key, "")).strip()
            if not text:
                issues.append({"kind": "不足", "what": f"考察 {no}) の必須（◎）「{what}」が無い"})
    if PLACEHOLDER in json.dumps(content, ensure_ascii=False):
        issues.append({"kind": "不足", "what": f"本文に {PLACEHOLDER} が残っている"})
    if not content.get("table3_max"):
        issues.append({"kind": "要確認", "what": "表3 の最大耐力と比べる理論値（Qmin・Qmean・Qy のどれか）が決まっていない"})
    return issues


def _set_fonts(doc, font: str = "IPAGothic") -> None:
    """日本語が PDF で正しく出るように、本文・見出しの欧文と東アジアのフォントを指定する（MS Mincho は無い環境を想定）。"""
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
    for st in doc.styles:
        if getattr(st, "type", None) not in (1,):   # 段落スタイル
            continue
        try:
            st.font.name = font
            rpr = st.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = rpr.makeelement(qn("w:rFonts"), {})
                rpr.append(rf)
            for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(k), font)
            for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rf.attrib.pop(qn(k), None)
            if st.name.startswith(("Heading", "Title")):
                st.font.color.rgb = RGBColor(0, 0, 0)
        except (AttributeError, ValueError):
            continue


def _page_number_footer(section) -> None:
    """フッター中央にページ番号（PAGE フィールド）。表紙を書式の PDF に差し替えても通しの番号になる。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    par = section.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def build(inp: Inputs, content: dict, outdir: Path, *, pdf: bool = True) -> dict:
    """docx（と PDF）を書く。issues が残る場合も書くが、結果に issues を返し完成扱いにしない。
    content.json の主なキー: cover・sections（blocks: text / eq / subheading / figure / table）・discussion（◎ と ☆）・
    table3_max・custom_tables（表1・表5 など独自の表）・after_sections（まとめ）・references（参考文献。最後に置く）・
    template_pages（表紙を書式の PDF から使う）・cover_overlay（表紙の欄に重ねて記入する文字。手書き不可のため）。"""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt

    outdir.mkdir(parents=True, exist_ok=True)
    issues = validate(content)
    figs = figures(inp, outdir / "figures")
    tabs = tables(inp)
    if content.get("table3_max") in ("Qmin", "Qmean", "Qy") and "max" in inp.events:
        c = theory(measured_params(inp.materials))
        sym = content["table3_max"]
        tq = inp.events["max"]["P_kN"] / 2
        label = content.get("table3_max_label", f"最大耐力（理論値は {sym}）")
        tabs["表3"]["rows"].append([label, f"{tq:.2f}", f"{c[sym] / 1000:.2f}", f"{tq / (c[sym] / 1000):.2f}"])
    for k, t in (content.get("custom_tables") or {}).items():
        tabs[k] = t
    for k, t in tabs.items():
        if any(PLACEHOLDER in str(v) for row in t["rows"] for v in row):
            issues.append({"kind": "不足", "what": f"{k} に入力の不足による {PLACEHOLDER} がある"})
    doc = Document()
    _set_fonts(doc)
    st = doc.styles["Normal"]
    st.font.size = Pt(10.5)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(22)
    sec.top_margin = sec.bottom_margin = Mm(20)
    _page_number_footer(sec)
    cover = content.get("cover", {})
    doc.add_heading(content.get("title", "2026 構造実験報告書"), level=0)
    for k in COVER_FIELDS:
        doc.add_paragraph(f"{k}：{cover.get(k, '')}")
    doc.add_page_break()
    if content.get("title_body"):
        tp = doc.add_paragraph()
        tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = tp.add_run(content["title_body"])
        r.bold = True
        r.font.size = Pt(13)
        if content.get("author_line"):
            ap = doc.add_paragraph(content["author_line"])
            ap.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fig_no, tab_no = [0], [0]

    def add_figure(key):
        fig_no[0] += 1
        doc.add_picture(str(figs[key]["path"]), width=Mm(150))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph(f"図-{fig_no[0]}　{figs[key]['caption']}")   # 図の題は図の下
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figs[key]["no"] = fig_no[0]

    def add_table(key):
        t = tabs[key]
        tab_no[0] += 1
        cap = doc.add_paragraph(f"{key.replace('表', '表-')}　{t['title']}")  # 書式の表番号。表の題は表の上（書式の例）
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tbl = doc.add_table(rows=1 + len(t["rows"]), cols=len(t["head"]))
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        ncol = len(t["head"])
        w0 = 62 if ncol <= 5 else 40
        widths = t.get("widths") or ([w0] + [(160 - w0) / (ncol - 1)] * (ncol - 1) if ncol > 1 else [160])
        for row in tbl.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Mm(w)
        for col, w in zip(tbl.columns, widths):
            col.width = Mm(w)
        for i, h in enumerate(t["head"]):
            cell = tbl.rows[0].cells[i]
            cell.text = h
            for par in cell.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.bold = True
                    run.font.size = Pt(9)
        for r, row in enumerate(t["rows"], start=1):
            for i, v in enumerate(row):
                cell = tbl.rows[r].cells[i]
                cell.text = str(v)
                for par in cell.paragraphs:
                    if i > 0 and not t.get("align_left"):
                        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in par.runs:
                        run.font.size = Pt(9)
        cap.paragraph_format.keep_with_next = True
        for row in tbl.rows[:-1]:
            for cell in row.cells:
                for par in cell.paragraphs:
                    par.paragraph_format.keep_with_next = True
        if t.get("note"):
            n = doc.add_paragraph(t["note"])
            for run in n.runs:
                run.font.size = Pt(8.5)
        else:
            doc.add_paragraph("")

    def add_blocks(blocks):
        for block in blocks:
            if "text" in block:
                doc.add_paragraph(block["text"])
            elif "eq" in block:
                ep = doc.add_paragraph(block["eq"])
                ep.paragraph_format.left_indent = Mm(8)
            elif "subheading" in block:
                doc.add_heading(block["subheading"], level=2)
            elif "figure" in block:
                add_figure(block["figure"])
            elif "table" in block:
                add_table(block["table"])
            elif "page_break" in block:
                doc.add_page_break()

    for s in content.get("sections", []):
        doc.add_heading(s["heading"], level=1)
        add_blocks(s.get("blocks", []))
    doc.add_heading(content.get("discussion_heading", "Ⅴ　考察"), level=1)
    disc = content.get("discussion", {})
    titles = content.get("discussion_titles", {})
    for no, items in MANDATORY.items():
        doc.add_heading(titles.get(no, f"{no})"), level=2)
        for key in list(items) + [k for k in disc.get(no, {}) if k not in items]:
            v = disc.get(no, {}).get(key)
            if not v:
                continue
            label = v.get("label") if isinstance(v, dict) else None
            if label is None:
                label = f"◎ {items[key].split('（')[0]}" if key in items else None
            if label:
                lp = doc.add_paragraph()
                lr = lp.add_run(label)
                lr.bold = True
            if isinstance(v, dict):
                add_blocks(v.get("blocks", []))
            else:
                doc.add_paragraph(str(v))
    for s in content.get("after_sections", []):
        doc.add_heading(s["heading"], level=1)
        add_blocks(s.get("blocks", []))
    if content.get("references"):
        doc.add_heading("参考文献", level=1)                      # 参考文献は最後（講義の指示）
        for i, ref in enumerate(content["references"], start=1):
            doc.add_paragraph(f"{i}) {ref}")
    docx_path = outdir / "report.docx"
    doc.save(docx_path)
    res = {"docx": docx_path, "issues": issues, "figures": fig_no[0], "tables": tab_no[0]}
    if pdf:
        import shutil
        import subprocess
        import tempfile
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        with tempfile.TemporaryDirectory() as prof:
            subprocess.run([soffice, f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf",
                            "--outdir", str(outdir), str(docx_path)], check=True, capture_output=True, timeout=300)
        res["pdf"] = outdir / "report.pdf"
        tpl = content.get("template_pages")   # 例: {"pdf": "work/inputs/…報告書.pdf", "pages": [1]}（表紙を書式の PDF から使う場合）
        if tpl:
            import pymupdf as fitz
            src = fitz.open(tpl["pdf"])
            body = fitz.open(res["pdf"])
            out = fitz.open()
            for n in tpl.get("pages", [1]):
                out.insert_pdf(src, from_page=n - 1, to_page=n - 1)
            out.insert_pdf(body, from_page=1 if tpl.get("replace_cover", True) else 0)
            overlay = content.get("cover_overlay") or []
            fontfile = _jp_font_file()
            if overlay and fontfile:
                for o in overlay:
                    out[o.get("page", 1) - 1].insert_text((o["x"], o["y"]), o["text"], fontsize=o.get("size", 11),
                                                          fontname="jp", fontfile=fontfile)
            out.save(outdir / "report_with_template.pdf")
            res["pdf"] = outdir / "report_with_template.pdf"
            if not overlay or not fontfile:
                issues.append({"kind": "要確認", "what": "書式の PDF の表紙の欄（学生番号・氏名・日付）が未記入（手書き不可の指示あり）"})
    return res
