#!/usr/bin/env python3
"""Excel・レポートの作成と検査（RC 梁の破壊実験。AI は呼ばない・無料）。手順は docs/DELIVERABLES.md。

  python3 scripts/deliver.py check-inputs --inputs DIR                 入力（CSV・材料強度・データ・readme）の検査（読み取りだけ）
  python3 scripts/deliver.py excel  --inputs DIR --out DIR --approved   Excel を作り、LibreOffice で再計算して照合
  python3 scripts/deliver.py report --inputs DIR --content JSON --out DIR --approved   レポート（docx・PDF）を組み立てて検査

excel・report は提出物を作るので、利用者が作成を許可した後だけ --approved を付けて実行する。
出力（実際の値・個人情報を含む）は work/ に置き、コミットしない。
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))


def check_inputs(inputs: Path) -> dict:
    from structlab.deliver import rc_beam as R
    inp = R.load_inputs(inputs)
    summary = R.check(inp)
    t = R.theory(R.DESIGN)
    design_check = {k: {"printed": v, "computed": t[k], "diff_pct": round((t[k] / v - 1) * 100, 3)} for k, v in R.PRINTED.items()}
    return {"summary": summary, "materials": inp.materials, "events": inp.events, "mapping": inp.mapping,
            "design_formula_check": design_check, "issues": inp.issues}


def excel(inputs: Path, out: Path) -> dict:
    import openpyxl

    from structlab.deliver import excel as X
    from structlab.deliver import rc_beam as R
    from structlab.checks.formula import close_enough
    inp = R.load_inputs(inputs)
    R.check(inp)
    src = X.build_workbook(inp, out / "src" / "実験データ解析.xlsx")
    rc = X.recalc_with_libreoffice(src, out)
    wb = openpyxl.load_workbook(rc, data_only=True)
    exp = X.expected_values(inp)
    got = {row[0].split()[-1]: (row[1], row[2]) for row in wb["表4"].iter_rows(min_row=2, values_only=True)}
    mism = [k for k, (fq, cq) in exp.items() if not (close_enough(got[k][0], fq, rel=1e-9) and close_enough(got[k][1], cq, rel=1e-9))]
    errors = [f"{ws.title}!{c.coordinate}" for ws in wb for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#")]
    return {"xlsx": str(rc), "formula_vs_python_mismatch": mism, "error_cells": errors[:20], "issues": inp.issues}


def report(inputs: Path, content: Path, out: Path) -> dict:
    from structlab.deliver import rc_beam as R
    from structlab.deliver import report as RP
    inp = R.load_inputs(inputs)
    R.check(inp)
    res = RP.build(inp, json.loads(content.read_text(encoding="utf-8")), out)
    res["issues"] = inp.issues + res["issues"]
    return {k: str(v) if isinstance(v, Path) else v for k, v in res.items()}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("action", choices=["check-inputs", "excel", "report"])
    ap.add_argument("--inputs", default=str(ROOT / "work" / "inputs"))
    ap.add_argument("--out", default=str(ROOT / "work" / "deliver"))
    ap.add_argument("--content")
    ap.add_argument("--approved", action="store_true", help="利用者が提出物の作成を許可した後だけ付ける")
    a = ap.parse_args()
    if a.action != "check-inputs" and not a.approved:
        raise SystemExit("提出物（Excel・レポート）の作成は利用者の許可の後だけ（--approved）")
    if a.action == "check-inputs":
        r = check_inputs(Path(a.inputs))
    elif a.action == "excel":
        r = excel(Path(a.inputs), Path(a.out) / "excel")
    else:
        if not a.content:
            raise SystemExit("--content（本文の JSON）が必要です")
        r = report(Path(a.inputs), Path(a.content), Path(a.out) / "report")
    print(json.dumps(r, ensure_ascii=False, indent=1, default=str))
    blocking = [i for i in r.get("issues", []) if i["kind"] in ("不足", "矛盾")]
    return 1 if blocking or r.get("formula_vs_python_mismatch") or r.get("error_cells") else 0


if __name__ == "__main__":
    raise SystemExit(main())
